"""Unit tests for ERP interoperability and export adapters (Section 12.0)."""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from api.database import SyncSessionLocal
from api.main import app
from api.security.jwt_auth import create_access_token
from api.services.erp import (
    generate_erp_export_payload,
    resolve_field_lengths,
    serialize_to_csv,
    serialize_to_xlsx,
    truncate_clean,
)


@pytest.fixture
def auth_headers() -> dict[str, str]:
    token = create_access_token(sub="usr-oil-maker-1", org_id="OIL-1", role="MAKER")
    return {"Authorization": f"Bearer {token}"}


def test_truncate_clean() -> None:
    """Tests clean description truncation without mid-word cutting and stripping punctuation."""
    # Under limit
    assert truncate_clean("GATE VALVE", 20) == "GATE VALVE"
    assert truncate_clean("", 20) == ""

    # Clean word boundary truncation (avoids clipping 'API600' mid-token at 40 chars)
    long_desc = "VALVE GATE 100 NB CL150 WCB FLG RF API600 INDUSTRIAL"
    clipped_40 = truncate_clean(long_desc, 40)
    assert len(clipped_40) <= 40
    # Should not end with dangling comma, dash, slash, or mid-word
    assert not clipped_40.endswith(("-", "/", ",", "|", " "))
    assert clipped_40 == "VALVE GATE 100 NB CL150 WCB FLG RF"

    # Extreme length constraint
    clipped_10 = truncate_clean("LONGTEXTWITHOUTSPACES", 10)
    assert len(clipped_10) <= 10


def test_resolve_field_lengths() -> None:
    """Tests resolving target field lengths with org config and defaults."""
    # Defaults
    sap_lens = resolve_field_lengths(None, "sap")
    assert sap_lens["description"] == 40
    assert sap_lens["code"] == 18

    oracle_lens = resolve_field_lengths(None, "oracle")
    assert oracle_lens["description"] == 60
    assert oracle_lens["code"] == 20

    maximo_lens = resolve_field_lengths(None, "maximo")
    assert maximo_lens["description"] == 50
    assert maximo_lens["code"] == 16

    # Org override
    custom_profile = {"field_lengths": {"description": 35, "code": 14}}
    custom_lens = resolve_field_lengths(custom_profile, "sap")
    assert custom_lens["description"] == 35
    assert custom_lens["code"] == 14


def test_generate_erp_export_payload_sap() -> None:
    """Tests SAP export formatting with MATNR, MAKTX, and BISMT carrying national code."""
    with SyncSessionLocal() as db:
        payload = generate_erp_export_payload(db=db, org_id="OIL-1", target="sap")

        assert payload["target"] == "sap"
        assert payload["org_id"] == "OIL-1"
        assert "items" in payload
        assert "classifications" in payload
        assert payload["field_lengths"]["description"] == 40

        items = payload["items"]
        if items:
            first_item = items[0]
            assert "MATNR" in first_item
            assert "MAKTX" in first_item
            assert "BISMT" in first_item
            assert "MEINS" in first_item
            assert len(first_item["MAKTX"]) <= 40
            assert first_item["BISMT"].startswith("MPNG-")

        classifications = payload["classifications"]
        if classifications:
            first_c = classifications[0]
            assert "OBJEK" in first_c
            assert "KLART" in first_c
            assert "CLASS" in first_c
            assert "CHARACT" in first_c
            assert "VALUE" in first_c


def test_generate_erp_export_payload_oracle() -> None:
    """Tests Oracle export formatting with ITEM_NUMBER, DESCRIPTION, and CROSS_REFERENCE."""
    with SyncSessionLocal() as db:
        payload = generate_erp_export_payload(db=db, org_id="STEEL-1", target="oracle")

        assert payload["target"] == "oracle"
        items = payload["items"]
        if items:
            first = items[0]
            assert "ITEM_NUMBER" in first
            assert "DESCRIPTION" in first
            assert "CROSS_REFERENCE" in first
            assert "CROSS_REFERENCE_TYPE" in first
            assert first["CROSS_REFERENCE_TYPE"] == "NATIONAL_MATERIAL_CODE"
            assert len(first["DESCRIPTION"]) <= 60


def test_generate_erp_export_payload_maximo() -> None:
    """Tests Maximo export formatting with ITEMNUM, DESCRIPTION, and EXTERNAL_REF."""
    with SyncSessionLocal() as db:
        payload = generate_erp_export_payload(db=db, org_id="MINE-1", target="maximo")

        assert payload["target"] == "maximo"
        items = payload["items"]
        if items:
            first = items[0]
            assert "ITEMNUM" in first
            assert "DESCRIPTION" in first
            assert "EXTERNAL_REF" in first
            assert "COMMODITY" in first
            assert len(first["DESCRIPTION"]) <= 50


def test_serialize_csv_and_xlsx() -> None:
    """Tests CSV and multi-sheet XLSX serialization helpers."""
    sample_payload = {
        "items": [
            {"MATNR": "1002341", "MAKTX": "VALVE GATE 100 NB", "BISMT": "MPNG-0004821937-3", "MEINS": "EA"},
        ],
        "classifications": [
            {"OBJEK": "1002341", "KLART": "001", "CLASS": "VALVES_GATE", "CHARACT": "SIZE_NB", "VALUE": "100"},
        ],
    }

    # CSV
    csv_out = serialize_to_csv(sample_payload["items"])
    assert "MATNR,MAKTX,BISMT,MEINS" in csv_out
    assert "1002341,VALVE GATE 100 NB,MPNG-0004821937-3,EA" in csv_out

    # XLSX
    xlsx_bytes = serialize_to_xlsx(sample_payload)
    assert len(xlsx_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Item_Master" in wb.sheetnames
    assert "Classifications" in wb.sheetnames


def test_erp_export_api_endpoints(auth_headers: dict[str, str]) -> None:
    """Tests GET /erp/export API endpoint across json, csv, and xlsx formats."""
    client = TestClient(app)

    # 1. JSON Export
    res_json = client.get("/erp/export?target=sap&format=json", headers=auth_headers)
    assert res_json.status_code == 200
    data = res_json.json()
    assert data["target"] == "sap"
    assert "items" in data
    assert "classifications" in data

    # 2. CSV Export
    res_csv = client.get("/erp/export?target=sap&format=csv&feed=items", headers=auth_headers)
    assert res_csv.status_code == 200
    assert res_csv.headers["content-type"].startswith("text/csv")
    assert "MATNR" in res_csv.text

    # 3. XLSX Export
    res_xlsx = client.get("/erp/export?target=oracle&format=xlsx", headers=auth_headers)
    assert res_xlsx.status_code == 200
    assert "spreadsheetml.sheet" in res_xlsx.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    assert "Item_Master" in wb.sheetnames
